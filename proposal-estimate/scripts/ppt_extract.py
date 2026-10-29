"""PPT 글자 추출 — 쪽마다 `===== p.N =====` 머리를 달아 텍스트로 뽑는다.

노트의 근거 표기(p.N)와 PM 페르소나에게 줄 입력을 같은 파일에서 만든다.

사용:
    uv run --with python-pptx python ppt_extract.py 제안서.pptx -o 제안서.txt
    uv run --with python-pptx python ppt_extract.py 제안서.pptx -o pm_input.txt \
        --exclude 29-38 --drop-money --drop-effort

--drop-money / --drop-effort 로 거른 줄은 `<출력>.dropped.txt` 에 따로 남긴다.
세션이 무엇이 빠졌는지 확인하는 용도이고, PM 페르소나에게는 주지 않는다.
"""

import argparse
import re
import sys
from pathlib import Path

# 괄호 안 금액: (1,200천원) (3억원) 처럼 문장 속에 끼어 있는 것은 그 조각만 지운다
MONEY_PAREN = re.compile(r"\(\s*[₩\\]?\s*\d[\d,.]*\s*(?:원|천원|만원|백만원|천만원|억원|억)\s*\)")
# 그 밖의 금액이 남은 줄은 통째로 거른다
MONEY = re.compile(
    r"(?:[₩\\]\s*\d)|(?:\d[\d,.]*\s*(?:원|천원|만원|백만원|천만원|억원|억)(?![가-힣]))|(?:KRW\s*\d)|(?:\d\s*KRW)",
)
# 고객 쪽 공수 초안(M/D, M/M, 인일)도 앵커가 되므로 따로 거를 수 있게 둔다
# 대소문자를 가린다: 「10mm」 같은 치수를 공수로 잘못 거르지 않게
EFFORT = re.compile(r"\d[\d,.]*\s*(?:[Mm]/[Dd]|M/M|MD|MM|맨데이|[Mm]an-?[Dd]ays?|인일|인월)(?![A-Za-z])")


def parse_pages(spec):
    # "1-28,39-44" 를 쪽 번호 집합으로
    if not spec:
        return None
    pages = set()
    for part in spec.split(","):
        part = part.strip()
        if not part:
            continue
        if "-" in part:
            a, b = part.split("-", 1)
            pages.update(range(int(a), int(b) + 1))
        else:
            pages.add(int(part))
    return pages


def slide_lines(slide, with_notes=False):
    out = []

    def walk(shapes):
        for sh in shapes:
            if sh.shape_type == 6:  # 그룹 도형
                walk(sh.shapes)
                continue
            if getattr(sh, "has_text_frame", False) and sh.has_text_frame:
                t = sh.text_frame.text.strip()
                if t:
                    out.extend(t.splitlines())
            if getattr(sh, "has_table", False) and sh.has_table:
                for r in sh.table.rows:
                    out.append(" | ".join(c.text.strip().replace("\n", " ") for c in r.cells))

    walk(slide.shapes)
    if with_notes and slide.has_notes_slide:
        note = slide.notes_slide.notes_text_frame.text.strip()
        if note:
            out.append("[발표자 노트]")
            out.extend(note.splitlines())
    return out


def filter_lines(lines, drop_money=False, drop_effort=False):
    # 남길 줄과 거른 줄을 나눈다
    kept, dropped = [], []
    for line in lines:
        s = line
        if drop_money:
            s = MONEY_PAREN.sub("", s)
            if MONEY.search(s):
                dropped.append(line)
                continue
        if drop_effort and EFFORT.search(s):
            dropped.append(line)
            continue
        if s != line:
            dropped.append(line + "    (괄호 금액만 지움)")
        kept.append(s.rstrip())
    return kept, dropped


def extract(pptx_path, pages=None, exclude=None, drop_money=False, drop_effort=False, with_notes=False):
    from pptx import Presentation

    prs = Presentation(str(pptx_path))
    text, dropped = [], []
    for i, slide in enumerate(prs.slides, 1):
        if pages and i not in pages:
            continue
        if exclude and i in exclude:
            continue
        kept, gone = filter_lines(slide_lines(slide, with_notes), drop_money, drop_effort)
        text.append(f"===== p.{i} =====")
        text.extend(kept)
        for g in gone:
            dropped.append(f"p.{i}: {g}")
    return "\n".join(text) + "\n", dropped, len(prs.slides)


def main(argv=None):
    ap = argparse.ArgumentParser(description="PPT 글자를 쪽 번호와 함께 뽑는다")
    ap.add_argument("pptx")
    ap.add_argument("-o", "--out", required=True, help="출력 텍스트 파일")
    ap.add_argument("--pages", help="뽑을 쪽 (예: 1-28,39-44)")
    ap.add_argument("--exclude", help="뺄 쪽 (예: 29-38)")
    ap.add_argument("--drop-money", action="store_true", help="금액이 적힌 줄을 거른다")
    ap.add_argument("--drop-effort", action="store_true", help="M/D · M/M 같은 공수 숫자가 적힌 줄을 거른다")
    ap.add_argument("--notes", action="store_true", help="발표자 노트도 넣는다")
    a = ap.parse_args(argv)

    body, dropped, total = extract(
        a.pptx, parse_pages(a.pages), parse_pages(a.exclude), a.drop_money, a.drop_effort, a.notes
    )
    out = Path(a.out)
    out.write_text(body, encoding="utf-8", newline="\n")
    side = out.with_name(out.name + ".dropped.txt")
    if dropped:
        side.write_text("\n".join(dropped) + "\n", encoding="utf-8", newline="\n")
    elif side.exists():
        side.unlink()
    pages_out = body.count("===== p.")
    print(f"OK: 전체 {total}쪽 중 {pages_out}쪽, 거른 줄 {len(dropped)}개 -> {out}")
    if dropped:
        print(f"거른 줄 목록: {side}", file=sys.stderr)
    return 0


if __name__ == "__main__":
    sys.exit(main())
